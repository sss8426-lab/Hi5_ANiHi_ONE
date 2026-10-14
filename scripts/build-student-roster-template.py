"""학생 관리 › 엑셀 업로드용 '학생 명단 기본 양식'을 만든다 (출석부 기본 양식과 같은 모양).

  python scripts/build-student-roster-template.py

결과: public/admissions-web/renderer/templates/student-roster-template.xlsx
- 첫 시트 '학생명단'이 실제로 올라가는 시트(업로드는 처음 보이는 시트의 '이름' 줄 아래를 읽음).
- 칸 이름은 student-grades.js 의 HEADERS 와 같아야 한다. 전공 계열·학년·연도는 업로드 화면에서 고르므로 넣지 않는다.
- 실기능력·희망 계열 목록은 app.js 의 SKILL_LEVELS · BASE_HOPE_TRACKS 와 같게 유지한다.
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.protection import SheetProtection

OUT = Path(__file__).resolve().parent.parent / 'public/admissions-web/renderer/templates/student-roster-template.xlsx'
SKILL_LEVELS = ['A+', 'A0', 'A-', 'B+', 'B0', 'B-', 'C']
TRACKS = ['기초소양', '기초디자인', '발상과표현', '상황표현', '칸만화', '이미지보드', '포트폴리오']
COLUMNS = [  # (머리글, 폭, 필수)
    ('No', 5, False), ('이름', 12, True), ('학교', 16, False), ('내신', 8, False), ('실기능력', 10, False),
    ('희망 계열', 14, False), ('희망 대학', 26, False), ('메모', 40, False),
]
ROWS = 200
FONT = '맑은 고딕'
TITLE_FILL, HEAD_FILL, LABEL_FILL = PatternFill('solid', fgColor='1F4E78'), PatternFill('solid', fgColor='5B9BD5'), PatternFill('solid', fgColor='D9EAF7')
NEED_FILL, NO_FILL, NOTE_FILL = PatternFill('solid', fgColor='FFF2CC'), PatternFill('solid', fgColor='F2F2F2'), PatternFill('solid', fgColor='FFF9E6')
thin = Side(style='thin', color='000000')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
LAST = chr(ord('A') + len(COLUMNS) - 1)

EXAMPLES = [
    ('김예시', '예시고', 2.3, 'A0', '상황표현', '한국대 만화애니', '주 2회 · 3월 입시 상담 예정'),
    ('이샘플', '샘플여고', 3.1, 'B+', '칸만화', '서울대 디자인', ''),
    ('박연습', '연습고', '', 'B0', '기초디자인', '', '내신은 성적표 받으면 입력'),
    ('최가상', '가상정보산업고', 4.2, 'C', '포트폴리오', '경기대 · 강원대', '특성화고'),
    ('정테스트', '테스트고', 1.8, 'A+', '발상과표현', '', ''),
]


def table_sheet(ws, title, note, examples=()):
    ws.sheet_view.showGridLines = False
    for i, (_, width, _) in enumerate(COLUMNS):
        ws.column_dimensions[chr(ord('A') + i)].width = width
    ws.merge_cells(f'A1:{LAST}1')
    ws['A1'] = title
    ws['A1'].font, ws['A1'].fill, ws['A1'].alignment = Font(name=FONT, size=15, bold=True, color='FFFFFF'), TITLE_FILL, CENTER
    ws.row_dimensions[1].height = 30
    ws['A2'] = '안내'
    ws['A2'].font, ws['A2'].fill, ws['A2'].alignment, ws['A2'].border = Font(name=FONT, size=9, bold=True, color='17365D'), LABEL_FILL, CENTER, BOX
    ws.merge_cells(f'B2:{LAST}2')
    ws['B2'] = note
    ws['B2'].font, ws['B2'].fill, ws['B2'].alignment = Font(name=FONT, size=9, color='17365D'), NOTE_FILL, LEFT
    for col in range(2, len(COLUMNS) + 1):
        ws.cell(row=2, column=col).border = BOX
    ws.row_dimensions[2].height = 30
    for i, (name, _, need) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=3, column=i, value=name)  # 업로드가 이 이름으로 칸을 찾음
        cell.font, cell.fill, cell.alignment, cell.border = Font(name=FONT, size=10, bold=True, color='FFFFFF'), HEAD_FILL, CENTER, BOX
    ws.row_dimensions[3].height = 22
    for r in range(ROWS):
        row = 4 + r
        ws.row_dimensions[row].height = 18
        for i, (name, _, need) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row, column=i)
            cell.border = BOX
            cell.font = Font(name=FONT, size=10)
            if name == 'No':
                cell.value, cell.fill, cell.alignment = r + 1, NO_FILL, CENTER
                continue
            cell.protection = Protection(locked=False)
            cell.alignment = LEFT if name in ('희망 대학', '메모') else CENTER
            if need:
                cell.fill = NEED_FILL
            if name == '내신':
                cell.number_format = '0.00'
        if r < len(examples):
            for i, value in enumerate(examples[r], start=2):
                if value != '':
                    ws.cell(row=row, column=i, value=value)
    end = 3 + ROWS
    gpa = DataValidation(type='decimal', operator='between', formula1='1', formula2='9', allow_blank=True, showErrorMessage=True,
                         errorTitle='내신', error='내신은 1~9 사이 숫자로 적어 주세요. 예: 2.35', promptTitle='내신', prompt='평균 등급 숫자. 모르면 비워 두세요.', showInputMessage=True)
    skill = DataValidation(type='list', formula1='"' + ','.join(SKILL_LEVELS) + '"', allow_blank=True, showErrorMessage=True,
                           errorTitle='실기능력', error='목록에서 골라 주세요 (A+ ~ C).', promptTitle='실기능력', prompt='칸 오른쪽 ▼를 눌러 고르세요.', showInputMessage=True)
    track = DataValidation(type='list', formula1='"' + ','.join(TRACKS) + '"', allow_blank=True, showErrorMessage=False,
                           promptTitle='희망 계열', prompt='▼ 목록에서 고르거나 직접 적어도 됩니다 (예: 웹툰).', showInputMessage=True)
    for dv, col in ((gpa, 'D'), (skill, 'E'), (track, 'F')):
        ws.add_data_validation(dv)
        dv.add(f'{col}4:{col}{end}')
    ws.freeze_panes = 'C4'
    ws.print_title_rows = '3:3'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f'A1:{LAST}{end}'
    ws.protection = SheetProtection(sheet=True, formatColumns=False, formatRows=False, insertRows=False, deleteRows=False,
                                    sort=False, autoFilter=False, selectLockedCells=False, selectUnlockedCells=False)


def guide_sheet(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 4
    ws.column_dimensions['B'].width = 100
    ws.merge_cells('A1:B1')
    ws['A1'] = '학생 명단 기본 양식 · 사용안내'
    ws['A1'].font, ws['A1'].fill, ws['A1'].alignment = Font(name=FONT, size=15, bold=True, color='FFFFFF'), TITLE_FILL, CENTER
    ws.row_dimensions[1].height = 30
    lines = [
        ('1', "'학생명단' 시트에 한 줄에 학생 한 명씩 적습니다. 노란 '이름' 칸은 꼭 적고, 나머지는 아는 것만 적어도 됩니다."),
        ('2', '내신은 평균 등급 숫자(예: 2.35), 실기능력과 희망 계열은 칸 오른쪽 ▼ 목록에서 고릅니다. 희망 계열은 목록에 없으면 직접 적어도 됩니다.'),
        ('3', '전공 계열(디자인·애니), 학년(고1~고3), 몇 년도 학년인지는 엑셀에 적지 않습니다. 올릴 때 화면에서 고릅니다.'),
        ('4', '그래서 디자인·애니, 학년이 다른 학생은 파일을 나눠서(또는 같은 파일을 여러 번) 올립니다.'),
        ('5', "DATA CORE › 상담용 › 대학 합격 로드맵 › 학생 관리 › 학년 탭 › [엑셀 업로드] → 전공 계열·학년·연도 고르기 → 이 파일 선택 → [올리기]."),
        ('6', "'작성예시' 시트는 보기용이라 올라가지 않습니다. 올라가는 것은 맨 앞 '학생명단' 시트뿐입니다."),
        ('7', '이미 있는 학생을 다시 올리면 같은 이름 학생이 하나 더 생깁니다. 처음 등록할 때만 쓰고, 이후 수정은 학생 화면에서 해 주세요.'),
        ('8', "성적은 이 양식 대신 학년 탭의 [성적표 사진 넣기]로 성적표 사진을 넣으면 자동으로 들어갑니다."),
        ('9', '줄이 모자라면 아래에 줄을 추가해도 됩니다. 칸 이름(3번째 줄)은 바꾸지 마세요.'),
    ]
    for i, (num, text) in enumerate(lines, start=3):
        a, b = ws.cell(row=i, column=1, value=num), ws.cell(row=i, column=2, value=text)
        a.font, a.alignment, a.fill = Font(name=FONT, size=10, bold=True, color='FFFFFF'), CENTER, PatternFill('solid', fgColor='174F44')
        b.font, b.alignment = Font(name=FONT, size=11), LEFT
        ws.row_dimensions[i].height = 34
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f'A1:B{2 + len(lines)}'
    ws.protection = SheetProtection(sheet=True, selectLockedCells=False, selectUnlockedCells=False)


def main():
    wb = Workbook()
    roster = wb.active
    roster.title = '학생명단'
    table_sheet(roster, 'CORE 학생 명단 입력', '전공 계열·학년·연도는 [엑셀 업로드] 화면에서 고릅니다 · 이름은 꼭, 나머지는 아는 것만 · 실기능력·희망 계열은 ▼ 목록에서')
    table_sheet(wb.create_sheet('작성예시'), '작성예시 (이 시트는 올라가지 않습니다)', '예시 학생은 모두 가상 인물입니다 · 실제 입력은 맨 앞 학생명단 시트에', EXAMPLES)
    guide_sheet(wb.create_sheet('사용안내'))
    wb.active = 0
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
